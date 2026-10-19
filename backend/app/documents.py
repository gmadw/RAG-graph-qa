"""文档解析与分块：支持 PDF / Word / Markdown / TXT / 图片（扫描件走 OCR）。"""
import asyncio
import re
from pathlib import Path

from . import db

# 可直接 OCR 的图片格式
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}


# ---------------- 纯文本类 ----------------


def extract_text(path: Path) -> str:
    """不涉及 OCR 的同步解析（docx / md / txt）。PDF 请用 extract_pdf_pages。"""
    suffix = path.suffix.lower()
    if suffix in (".docx",):
        import docx

        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    if suffix in (".md", ".markdown", ".txt", ""):
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"不支持的文件类型：{suffix}（支持 pdf/docx/md/txt/png/jpg）")


def pdf_page_texts(path: Path) -> list:
    """用 pypdf 逐页抽取文字层。扫描件没有文字层，返回的基本是空串。"""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return pages


# ---------------- 图片预处理与 OCR ----------------


def _prepare_image(raw: bytes, max_side: int = 2000) -> bytes:
    """统一转 JPEG 并限制长边，避免手机拍照原图过大导致请求超限或费用暴涨。"""
    from io import BytesIO

    from PIL import Image

    try:
        img = Image.open(BytesIO(raw))
        img = img.convert("RGB")
        w, h = img.size
        if max(w, h) > max_side:
            scale = max_side / max(w, h)
            img = img.resize((int(w * scale), int(h * scale)), Image.LANCZOS)
        buf = BytesIO()
        img.save(buf, format="JPEG", quality=88, optimize=True)
        return buf.getvalue()
    except Exception:
        return raw  # 解码失败就原样送，让上层拿到明确的接口报错


def _render_pdf_page(path: Path, index: int, dpi: int) -> bytes:
    """把 PDF 指定页渲染成 PNG 字节。"""
    import pymupdf  # PyMuPDF ≥1.24；旧版的 fitz 已弃用

    doc = pymupdf.open(str(path))
    try:
        page = doc[index]
        pix = page.get_pixmap(dpi=dpi)
        return pix.tobytes("png")
    finally:
        doc.close()


async def _ocr_pages(path: Path, indexes: list, dpi: int, on_progress=None) -> tuple:
    """并发 OCR 指定页，返回 ({页序号: 文本}, {页序号: 错误})。单页失败不中断整份文档。"""
    from .providers import ProviderError, ocr_image

    out, errors = {}, {}
    sem = asyncio.Semaphore(3)  # 免费档常有并发限制，3 路比较稳
    done = 0
    total = len(indexes)
    lock = asyncio.Lock()

    async def one(idx: int):
        nonlocal done
        async with sem:
            try:
                png = await asyncio.to_thread(_render_pdf_page, path, idx, dpi)
                text = await ocr_image(_prepare_image(png))
                async with lock:
                    out[idx] = text
            except ProviderError as e:
                async with lock:
                    errors[idx] = str(e)
            except Exception as e:  # 渲染失败等
                async with lock:
                    errors[idx] = f"{type(e).__name__}: {e}"
            async with lock:
                done += 1
                if on_progress:
                    on_progress(done, total)

    await asyncio.gather(*(one(i) for i in indexes))
    if errors and not out:
        first = next(iter(errors.values()))
        raise RuntimeError(f"OCR 全部失败：{first}")
    return out, errors


def _detect_needs_ocr(page_texts: list, min_chars: int) -> list:
    """返回需要 OCR 的页序号：文字层过少的页。"""
    return [i for i, t in enumerate(page_texts) if len(re.sub(r"\s", "", t or "")) < min_chars]


# ---------------- 统一入口 ----------------


async def extract_any(path: Path, cfg: dict, on_progress=None) -> tuple:
    """解析任意受支持文件，返回 (文本, 说明信息)。

    PDF 策略：先抽文字层；文字过少的页视为扫描页，在启用 OCR 时逐页识别。
    这样文字型 PDF 零额外开销，混合型（部分页扫描）只 OCR 缺文字的页。
    """
    suffix = path.suffix.lower()
    ocr_cfg = cfg.get("ocr") or {}
    ocr_on = bool(ocr_cfg.get("enabled"))
    min_chars = int(ocr_cfg.get("min_chars") or 50)
    dpi = int(ocr_cfg.get("dpi") or 150)
    max_pages = int(ocr_cfg.get("max_pages") or 50)

    # ---- 图片：整张走 OCR ----
    if suffix in IMAGE_SUFFIXES:
        if not ocr_on:
            raise ValueError("这是图片文件，需要文字识别（OCR）。请在「模型设置」页启用 OCR 后重试")
        from .providers import ocr_image

        raw = await asyncio.to_thread(path.read_bytes)
        text = await ocr_image(_prepare_image(raw))
        if not text.strip():
            raise ValueError("OCR 未识别到文字，图片可能过于模糊或没有文字")
        return text, "图片 OCR"

    # ---- PDF ----
    if suffix == ".pdf":
        page_texts = await asyncio.to_thread(pdf_page_texts, path)
        need = _detect_needs_ocr(page_texts, min_chars)
        scanned = len(need) == len(page_texts) and len(page_texts) > 0

        if not ocr_on:
            if scanned:
                raise ValueError(
                    "这份 PDF 没有文字层（扫描件/图片型 PDF），无法直接提取文字。"
                    "请在「模型设置」页启用 OCR（推荐硅基流动 PaddleOCR-VL-1.5）后重新上传"
                )
            return "\n".join(page_texts), "PDF 文字层"

        target = need[:max_pages]
        skipped = len(need) - len(target)
        if not target:
            return "\n".join(page_texts), "PDF 文字层"

        ocr_map, errors = await _ocr_pages(path, target, dpi, on_progress)
        parts = []
        for i, t in enumerate(page_texts):
            if i in ocr_map:
                parts.append(ocr_map[i])
            else:
                parts.append(t)
        note = f"{'扫描件' if scanned else '混合型'} PDF：OCR {len(ocr_map)}/{len(page_texts)} 页"
        if errors:
            note += f"，{len(errors)} 页失败"
        if skipped:
            note += f"，超过 max_pages 跳过 {skipped} 页"
        return "\n".join(parts), note

    # ---- 其他文本类 ----
    return await asyncio.to_thread(extract_text, path), ""


# 中文句子结束符优先切分
_SENT_SPLIT = re.compile(r"(?<=[。！？!?；;\n])")
_PARA_SPLIT = re.compile(r"\n\s*\n")


def _split_long(text: str, target: int) -> list:
    """对超长段落做句子级累积，硬切时带 overlap。"""
    out, cur = [], ""
    for sent in _SENT_SPLIT.split(text):
        if not sent.strip():
            continue
        if len(cur) + len(sent) <= target:
            cur += sent
        else:
            if cur.strip():
                out.append(cur.strip())
            while len(sent) > target:
                out.append(sent[:target].strip())
                sent = sent[target - 60:]
            cur = sent
    if cur.strip():
        out.append(cur.strip())
    return out


def chunk_text(text: str, target: int = 500, overlap_reserved: bool = True) -> list:
    """段落聚合 + 句子级兜底的目标分块。返回字符串列表。"""
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    paras = [p.strip() for p in _PARA_SPLIT.split(text) if p.strip()]
    chunks, cur = [], ""
    for p in paras:
        if len(cur) + len(p) + 2 <= target:
            cur = f"{cur}\n\n{p}".strip()
        else:
            if len(p) <= target * 1.6:
                if cur:
                    chunks.append(cur)
                cur = p
            else:
                if cur:
                    chunks.append(cur)
                    cur = ""
                chunks.extend(_split_long(p, target))
    if cur:
        chunks.append(cur)
    return chunks


async def process_document(doc_id: int, build_graph: bool = None):
    """后台任务：解析（含 OCR）→ 分块 → 向量化入库 →（可选）图谱构建。"""
    from . import config, graph
    from .providers import embed_texts, ProviderError

    doc = db.get_document(doc_id)
    if not doc:
        return
    try:
        cfg = config.get_settings()
        if build_graph is None:
            build_graph = cfg.get("graph", {}).get("build_on_upload", True)
        db.update_document(doc_id, status="processing", stage="解析文档", error="")

        path = db.UPLOAD_DIR / f"{doc_id}_{doc['filename']}"

        def on_progress(done, total):
            db.update_document(doc_id, stage=f"OCR 识别 {done}/{total} 页")

        text, note = await extract_any(path, cfg, on_progress=on_progress)
        chunks = chunk_text(text)
        if not chunks:
            db.update_document(doc_id, status="error", stage="", error="未能从文档中提取到文本")
            return

        db.update_document(
            doc_id, stage=f"向量化 {len(chunks)} 个分块" + (f"（{note}）" if note else "")
        )
        vectors = await embed_texts(chunks)
        import numpy as np
        tag = config.embed_tag(cfg)
        items = [(i, c, np.asarray(v, dtype=np.float32)) for i, (c, v) in enumerate(zip(chunks, vectors))]
        db.insert_chunks(doc_id, tag, items)
        db.update_document(doc_id, status="processing", stage="向量化完成", chunk_count=len(chunks), embed_tag=tag)

        if build_graph:
            db.update_document(doc_id, stage="构建知识图谱")
            triple_count = await graph.build_graph_for_doc(doc_id)
            db.update_document(doc_id, triple_count=triple_count)

        db.update_document(doc_id, status="done", stage="")
    except ProviderError as e:
        db.update_document(doc_id, status="error", error=f"模型调用失败：{e}")
    except Exception as e:
        db.update_document(doc_id, status="error", error=str(e)[:500])
